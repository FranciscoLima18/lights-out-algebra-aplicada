"""Genera el informe de Lights Out en formato DOCX.

Uso:
    python generar_informe_docx.py
    python generar_informe_docx.py --salida informe_lights_out.docx
"""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt


TITLE = "Modelado algebraico y resolución computacional del juego Lights Out"
STUDENT = "Francisco Lima Grille"
ID_NUMBER = "5.569.014-5"
DATE = "4 de octubre de 2026"
COURSE = "Álgebra Aplicada"


def configure_document(document: Document) -> None:
    """Configura márgenes, fuente base y estilos nativos de Word."""
    section = document.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2.5)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)

    for style_name, size in (("Heading 1", 14), ("Heading 2", 12)):
        style = document.styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header_run = header.add_run(
        f"{TITLE} | {STUDENT} | C.I. {ID_NUMBER} | {DATE}"
    )
    header_run.font.name = "Times New Roman"
    header_run.font.size = Pt(9)


def add_cover(document: Document) -> None:
    """Agrega el encabezado formal de identificación del informe."""
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(24)
    run = paragraph.add_run(TITLE)
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(18)

    for label, value in (
        ("Estudiante", STUDENT),
        ("Cédula de identidad", ID_NUMBER),
        ("Fecha", DATE),
        ("Curso", COURSE),
        (
            "Repositorio de código",
            "https://github.com/FranciscoLima18/lights-out-algebra-aplicada",
        ),
    ):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run(f"{label}: ").bold = True
        paragraph.add_run(value)

    document.add_page_break()


def add_heading(document: Document, text: str, level: int = 1) -> None:
    document.add_heading(text, level=level)


def add_normal(document: Document, text: str = "") -> None:
    document.add_paragraph(text, style="Normal")


def add_equation(document: Document, text: str) -> None:
    """Agrega una ecuación como texto plano Unicode, no como LaTeX."""
    paragraph = document.add_paragraph(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.font.name = "Cambria Math"


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        document.add_paragraph(item, style="List Bullet")


def add_results_table(document: Document) -> None:
    document.add_paragraph(
        "Tabla 1. Casos cubiertos por las pruebas automatizadas.",
        style="Normal",
    )
    rows = [
        ("Caso", "Resultado observado"),
        ("Tablero 1×1 apagado", "Devuelve [0]"),
        ("Tablero 1×1 encendido", "Devuelve [1]"),
        ("Tablero 3×3 con solución conocida", "Coincide con el vector esperado"),
        (
            "Tableros 3×3, 4×4 y 5×5",
            "La aplicación de la solución apaga el tablero",
        ),
        ("Sistema con solución no única", "Devuelve una solución binaria válida"),
        ("Sistema inconsistente", "Lanza ValueError indicando falta de solución"),
        (
            "Entradas no válidas",
            "Lanza ValueError para matrices no cuadradas o no binarias",
        ),
    ]
    table = document.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row_index, row_data in enumerate(rows):
        for column_index, value in enumerate(row_data):
            cell = table.cell(row_index, column_index)
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index == 0:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True
    note = document.add_paragraph("Nota. Resultados obtenidos de las ejecuciones registradas en test_lights_out.py.")
    note.runs[0].italic = True


def build_report() -> Document:
    document = Document()
    configure_document(document)
    add_cover(document)

    add_heading(document, "Resumen")
    add_normal(
        document,
        "El presente informe estudia el juego Lights Out mediante un modelo "
        "algebraico sobre el cuerpo Z₂. El objetivo es justificar que una "
        "solución puede representarse por un vector binario de dimensión n², "
        "construir las ecuaciones asociadas a un tablero n × n y verificar una "
        "solución propuesta para el caso 3 × 3. La metodología combina el "
        "análisis de las propiedades de las pulsaciones con la construcción "
        "de una matriz de adyacencia y la eliminación gaussiana utilizando "
        "operaciones XOR (suma módulo 2). Se demuestra que presionar una misma "
        "celda dos veces cancela su efecto y que el orden de las pulsaciones "
        "no modifica el estado final. En el sistema de ejemplo, la asignación "
        "x₅ = x₉ = 1 y xᵢ = 0 para los restantes índices satisface las nueve "
        "ecuaciones. Además, la implementación en Python resuelve tableros "
        "cuadrados binarios, identifica sistemas inconsistentes y devuelve una "
        "solución cuando existe. Por tanto, Lights Out se formula como un "
        "sistema lineal finito cuya resolución computacional se realiza "
        "mediante escalonamiento sobre Z₂.",
    )

    add_heading(document, "1. Introducción")
    add_normal(
        document,
        "Lights Out es un juego sobre un tablero cuadrado de luces inicialmente "
        "encendidas o apagadas. Al presionar una celda se modifica su estado y "
        "el de sus vecinos ortogonales. El objetivo consiste en alcanzar el "
        "estado completamente apagado.",
    )
    add_normal(
        document,
        "La interacción local de cada pulsación genera dependencias entre las "
        "celdas que pueden describirse mediante álgebra lineal sobre Z₂. Esta "
        "formulación transforma la búsqueda de jugadas en la resolución de un "
        "sistema de ecuaciones lineales A v = b (módulo 2).",
    )
    add_normal(
        document,
        "El objetivo general de este trabajo es desarrollar y justificar el "
        "modelo binario para Lights Out y verificarlo en el sistema correspondiente "
        "a un tablero 3 × 3. Como objetivos específicos se busca: (i) analizar "
        "la repetición y el orden de las pulsaciones; (ii) justificar la "
        "representación vectorial; (iii) explicar la construcción de las "
        "ecuaciones; y (iv) verificar la solución indicada mediante análisis "
        "teórico y pruebas computacionales.",
    )

    add_heading(document, "2. Marco teórico")
    add_heading(document, "2.1. Operaciones en Z₂", level=2)
    add_normal(document, "El conjunto Z₂ = {0, 1} opera con suma módulo 2:")
    add_equation(document, "0 + 0 = 0    0 + 1 = 1    1 + 0 = 1    1 + 1 = 0")
    add_normal(
        document,
        "Sumar dos veces el mismo elemento produce cero, lo que permite combinar "
        "pulsaciones mediante la función lógica XOR.",
    )
    add_heading(document, "2.2. Efecto de una pulsación", level=2)
    add_normal(
        document,
        "Cada celda a(q,r) cambia su propio estado y el de sus vecinas "
        "ortogonales pertenecientes al tablero. Las celdas en bordes o esquinas "
        "poseen menos vecinos adyacentes, reduciendo el número de términos en "
        "sus ecuaciones.",
    )
    add_heading(document, "2.3. Conmutatividad e involución", level=2)
    add_normal(
        document,
        "Las pulsaciones actúan como operadores de suma vectorial sobre Z₂. La "
        "suma vectorial es conmutativa (u + v = v + u), por lo que alterar la "
        "secuencia de pulsaciones no cambia el resultado final. Además, presionar "
        "una celda dos veces equivale a sumar su efecto dos veces, anulando el "
        "cambio (1 + 1 = 0).",
    )
    add_heading(document, "2.4. Representación vectorial", level=2)
    add_normal(document, "Para un tablero n × n, se asocia el vector de solución:")
    add_equation(document, "v = (x₁, x₂, …, xₙ²) ∈ Z₂ⁿ²")
    add_normal(
        document,
        "donde xᵢ = 1 si la celda correspondiente se presiona y xᵢ = 0 en "
        "caso contrario.",
    )

    add_heading(document, "3. Metodología")
    add_bullets(
        document,
        [
            "Construcción del vector de estado b: se aplanó la matriz inicial n × n, ordenando las celdas fila por fila.",
            "Matriz de adyacencia A: se construyó A ∈ Z₂ⁿ²×ⁿ², indicando con 1 si presionar la columna j conmuta la celda en la fila i.",
            "Escalerización gaussiana en Z₂: se aplicaron transformaciones Fᵢ → Fᵢ ⊕ Fⱼ para resolver A v = b (módulo 2).",
            "Validación: se desarrollaron pruebas automatizadas para tableros de 1 × 1 a 5 × 5, sistemas inconsistentes y casos no únicos.",
        ],
    )
    add_normal(document, "El sistema se expresa como:")
    add_equation(document, "A v = b (módulo 2)")
    add_normal(
        document,
        "La igualdad significa que el efecto de las pulsaciones debe coincidir "
        "con el estado inicial. En la implementación, la suma binaria se realiza "
        "mediante XOR. Si aparece una fila de la forma 0 = 1, el sistema es "
        "inconsistente; si hay variables libres, se fijan en cero y se devuelve "
        "una solución particular.",
    )

    add_heading(document, "4. Resultados")
    add_heading(document, "4.1. Pregunta 1: repetición de pulsaciones", level=2)
    add_normal(
        document,
        "No es necesario presionar una luz más de una vez. Como e + e = 0 "
        "módulo 2, presionar una celda dos veces anula la acción. Toda solución "
        "puede representarse de modo que cada celda se pulsa cero o una vez.",
    )
    add_heading(document, "4.2. Pregunta 2: relevancia del orden", level=2)
    add_normal(
        document,
        "El orden de las pulsaciones no es relevante. Como Pᵢ + Pⱼ = Pⱼ + Pᵢ "
        "módulo 2, cualquier permutación de la secuencia de pulsaciones genera "
        "el mismo estado final.",
    )
    add_heading(document, "4.3. Pregunta 3: validez del modelo", level=2)
    add_normal(
        document,
        "El modelo vectorial es correcto porque: (1) la respuesta de cada celda "
        "es binaria, (2) la repetición se reduce módulo 2 y (3) la suma es "
        "conmutativa. La división entera i − 1 = n · q + r, con 0 ≤ r < n, "
        "establece la biyección entre los elementos del vector v y la matriz "
        "del tablero.",
    )
    add_heading(document, "4.4. Pregunta 4: construcción de las ecuaciones (3 × 3)", level=2)
    add_normal(document, "Numerando las posiciones de 1 a 9 por filas:")
    add_equation(document, "1  2  3\n4  5  6\n7  8  9")
    add_bullets(
        document,
        [
            "Esquinas: E₁: x₁ + x₂ + x₄ = 0; E₃: x₂ + x₃ + x₆ = 0; E₇: x₄ + x₇ + x₈ = 0; E₉: x₆ + x₈ + x₉ = 1.",
            "Bordes: E₂: x₁ + x₂ + x₃ + x₅ = 1; E₄: x₁ + x₄ + x₅ + x₇ = 1; E₆: x₃ + x₅ + x₆ + x₉ = 0; E₈: x₅ + x₇ + x₈ + x₉ = 0.",
            "Centro: E₅: x₂ + x₄ + x₅ + x₆ + x₈ = 1.",
        ],
    )
    add_heading(document, "4.5. Pregunta 5: verificación de la solución propuesta", level=2)
    add_normal(document, "Sustituyendo x₅ = x₉ = 1 y los demás xᵢ = 0:")
    add_equation(document, "E₁: 0,  E₂: 1,  E₃: 0,  E₄: 1,  E₅: 1,  E₆: 0,  E₇: 0,  E₈: 0,  E₉: 1")
    add_normal(
        document,
        "El vector obtenido (0, 1, 0, 1, 1, 0, 0, 0, 1) coincide con los "
        "términos independientes del sistema inicial, confirmando la validez "
        "de la solución.",
    )
    add_heading(document, "4.6. Resultados computacionales", level=2)
    add_results_table(document)

    add_heading(document, "5. Discusión")
    add_normal(
        document,
        "El análisis confirma que la estructura algebraica de Z₂ permite reducir "
        "cualquier secuencia de acciones a un vector binario estático. La matriz "
        "de adyacencia maneja las condiciones de borde sin requerir reglas "
        "adicionales.",
    )
    add_normal(
        document,
        "Desde la perspectiva computacional, la eliminación gaussiana mediante "
        "XOR resuelve el problema eficientemente. En sistemas con soluciones "
        "no únicas, fijar variables libres en cero otorga una solución particular válida.",
    )

    add_heading(document, "6. Conclusiones")
    add_normal(
        document,
        "Lights Out se modela exitosamente como un sistema lineal A v = b sobre "
        "Z₂. Las propiedades de conmutatividad e involución justifican la "
        "representación mediante un vector de dimensión n². La verificación "
        "teórica y computacional desarrollada en Python valida la consistencia "
        "del algoritmo y de las soluciones propuestas.",
    )

    add_heading(document, "Referencias")
    add_normal(
        document,
        "Universidad Católica del Uruguay. (s. f.). Consigna del proyecto "
        "Lights Out (Álgebra Aplicada) [Material de curso].",
    )
    add_normal(
        document,
        "Universidad Católica del Uruguay. (s. f.). Guía de escritura de "
        "informes académicos [Guía institucional].",
    )

    add_heading(document, "Anexo A. Código utilizado")
    add_bullets(
        document,
        [
            "lights_out.py: módulo principal con la función solve_lights_out.",
            "test_lights_out.py: suite de pruebas unitarias.",
        ],
    )
    return document


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--salida",
        type=Path,
        default=Path("informe_lights_out.docx"),
        help="Ruta del archivo DOCX que se generará.",
    )
    args = parser.parse_args()
    args.salida.parent.mkdir(parents=True, exist_ok=True)
    build_report().save(args.salida)
    print(f"Documento generado: {args.salida}")


if __name__ == "__main__":
    main()
